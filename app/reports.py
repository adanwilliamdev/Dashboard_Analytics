"""Geracao de relatorios (Excel e PDF). Funcoes puras: entram dados, saem bytes."""
import io
import re
from typing import Any, Dict, List

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app import analytics_core as core
from app.utils import local_now

BRL = '"R$" #,##0.00'
HEADER_FILL = PatternFill("solid", fgColor="07090C")
HEADER_FONT = Font(bold=True, color="34D399")
_EMOJI = re.compile("[\U00010000-\U0010ffff\u2600-\u27bf\ufe0f]")


def money(value) -> str:
    if value is None:
        return "-"
    return "R$ " + f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _plain(text: Any) -> str:
    """Texto seguro para o PDF (fonte Helvetica nao tem emoji)."""
    return _EMOJI.sub("", str(text)).strip()


# --------------------------------------------------------------------- Excel
def _set(ws, row: int, col: int, value, fmt: str = None):
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"  # texto literal: "=1+1" NAO vira formula (injecao no Excel)
    if fmt:
        cell.number_format = fmt
    return cell


def _write_table(ws, headers: List[str], rows: List[List[Any]], formats: Dict[int, str] = None):
    formats = formats or {}
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            _set(ws, r, c, value, formats.get(c))
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for c, h in enumerate(headers, start=1):
        longest = max([len(str(h))] + [len(str(r[c - 1])) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(c)].width = min(max(longest + 2, 10), 50)


def build_transactions_xlsx(rows: List[Dict[str, Any]]) -> bytes:
    """rows: dicts com id, date(datetime), description, category, type, amount."""
    wb = Workbook()
    df = core.normalize(pd.DataFrame(rows, columns=core.COLUMNS)) if rows else core.normalize(None)

    # Resumo
    ws = wb.active
    ws.title = "Resumo"
    s = core.summarize(df)
    items = [
        ("Receitas", s["total_income"], BRL),
        ("Despesas", s["total_expense"], BRL),
        ("Saldo", s["balance"], BRL),
        ("Transações", s["transaction_count"], "0"),
        ("Ticket médio", s["avg_transaction"], BRL),
        ("Maior despesa", s["largest_expense"], BRL),
        ("Taxa de poupança (%)", s["savings_rate"], "0.0"),
    ]
    _set(ws, 1, 1, "Indicador").font = HEADER_FONT
    _set(ws, 1, 2, "Valor").font = HEADER_FONT
    ws["A1"].fill = ws["B1"].fill = HEADER_FILL
    for i, (label, value, fmt) in enumerate(items, start=2):
        _set(ws, i, 1, label)
        _set(ws, i, 2, value if value is not None else "-", fmt)
    _set(ws, len(items) + 3, 1, "Gerado em " + local_now().strftime("%d/%m/%Y %H:%M"))
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 18

    # Transacoes
    ws_t = wb.create_sheet("Transações")
    table = [
        [r["id"], pd.Timestamp(r["date"]).to_pydatetime(), r["description"], r["category"],
         "Receita" if r["type"] == "income" else "Despesa", float(r["amount"])]
        for r in rows
    ]
    _write_table(ws_t, ["ID", "Data", "Descrição", "Categoria", "Tipo", "Valor"], table, {2: "dd/mm/yyyy", 6: BRL})

    # Categorias
    ws_c = wb.create_sheet("Categorias")
    cats = core.by_category(df)
    _write_table(
        ws_c, ["Categoria", "Tipo", "Valor", "Qtde", "% do tipo"],
        [[c["category"], "Receita" if c["type"] == "income" else "Despesa", c["amount"], c["count"], c["share_pct"]]
         for c in cats],
        {3: BRL, 5: "0.0"},
    )

    # Mensal
    ws_m = wb.create_sheet("Mensal")
    monthly = core.monthly(df, 120)
    _write_table(
        ws_m, ["Mês", "Receita", "Despesa", "Saldo", "Poupança (%)"],
        [[m["month"], m["income"], m["expense"], m["balance"],
          m["savings_rate"] if m["savings_rate"] is not None else "-"] for m in monthly],
        {2: BRL, 3: BRL, 4: BRL, 5: "0.0"},
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------- PDF
def _table(data: List[List[Any]], col_widths=None, align_right_from: int = 1) -> Table:
    t = Table(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#07090C")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#34D399")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F4")]),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D6D3D1")),
                ("ALIGN", (align_right_from, 1), (-1, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    return t


def build_report_pdf(overview: Dict[str, Any], username: str = "") -> bytes:
    """Relatorio executivo (A4) a partir do payload de /analytics/overview."""
    styles = getSampleStyleSheet()
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], textColor=colors.HexColor("#07090C"), spaceBefore=12)
    small = ParagraphStyle("small", parent=styles["BodyText"], fontSize=8, textColor=colors.HexColor("#57534E"))

    s = overview["summary"]
    ch = s["comparison"]["change_pct"]
    days = overview["days"]

    def delta(v):
        return "-" if v is None else f"{v:+.1f}%"

    story: List[Any] = [
        Paragraph("Relatório financeiro", styles["Title"]),
        Paragraph(
            _plain(f"Últimos {days} dias · gerado em {local_now():%d/%m/%Y %H:%M}"
                   + (f" · {username}" if username else "")),
            small,
        ),
        Spacer(1, 0.4 * cm),
        _table(
            [
                ["Indicador", "Valor", f"vs. {days} dias anteriores"],
                ["Receitas", money(s["total_income"]), delta(ch["income"])],
                ["Despesas", money(s["total_expense"]), delta(ch["expense"])],
                ["Saldo", money(s["balance"]), delta(ch["balance"])],
                ["Taxa de poupança", "-" if s["savings_rate"] is None else f"{s['savings_rate']:.1f}%", ""],
                ["Gasto médio por dia", money(s["avg_daily_expense"]), ""],
                ["Transações", str(s["transaction_count"]), delta(ch["transactions"])],
            ],
            [6 * cm, 5 * cm, 6 * cm],
        ),
    ]

    if overview.get("insights"):
        story.append(Paragraph("Insights", h2))
        for ins in overview["insights"]:
            story.append(Paragraph("• " + _plain(ins["message"]), styles["BodyText"]))

    expense_cats = [c for c in overview["by_category"] if c["type"] == "expense"][:10]
    if expense_cats:
        story.append(Paragraph("Despesas por categoria", h2))
        story.append(
            _table(
                [["Categoria", "Valor", "% das despesas"]]
                + [[_plain(c["category"]), money(c["amount"]), f"{c['share_pct']:.1f}%"] for c in expense_cats],
                [7 * cm, 5 * cm, 5 * cm],
            )
        )

    if overview.get("top_expenses"):
        story.append(Paragraph("Maiores despesas", h2))
        story.append(
            _table(
                [["Data", "Descrição", "Categoria", "Valor"]]
                + [[t["date"][:10], _plain(t["description"])[:32], _plain(t["category"]), money(t["amount"])]
                   for t in overview["top_expenses"]],
                [2.5 * cm, 6.5 * cm, 4 * cm, 4 * cm], align_right_from=3,
            )
        )

    if overview.get("budgets"):
        story.append(Paragraph("Orçamentos do mês", h2))
        story.append(
            _table(
                [["Categoria", "Limite", "Gasto", "Uso"]]
                + [[_plain(b["category"]), money(b["monthly_limit"]), money(b["spent"]), f"{b['pct_used']:.0f}%"]
                   for b in overview["budgets"]],
                [5 * cm, 4 * cm, 4 * cm, 4 * cm],
            )
        )

    fc = overview.get("forecast") or {}
    if fc.get("available"):
        story.append(Paragraph("Previsão (próximos 30 dias)", h2))
        story.append(
            Paragraph(
                f"Despesa projetada: <b>{money(fc['projected_expense'])}</b> "
                f"(faixa {money(fc['projected_expense_lower'])} a {money(fc['projected_expense_upper'])}). "
                f"Receita projetada: <b>{money(fc['projected_income'])}</b>. "
                f"Saldo projetado: <b>{money(fc['projected_balance'])}</b>.",
                styles["BodyText"],
            )
        )
        story.append(Paragraph("Estimativa estatística simples; não é garantia de resultado.", small))

    if overview.get("anomalies"):
        story.append(Paragraph("Gastos atípicos", h2))
        story.append(
            _table(
                [["Data", "Descrição", "Valor", "Típico"]]
                + [[a["date"], _plain(a["description"])[:32], money(a["amount"]), money(a["expected"])]
                   for a in overview["anomalies"][:8]],
                [2.5 * cm, 7 * cm, 4 * cm, 4 * cm], align_right_from=2,
            )
        )

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                      bottomMargin=2 * cm, title="Relatório financeiro").build(story)
    return buf.getvalue()
