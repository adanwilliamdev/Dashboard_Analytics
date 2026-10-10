import io

from openpyxl import load_workbook

from tests.conftest import register_and_login
from tests.test_transactions import add


def test_admin_by_env_and_regular_user_blocked(client):
    admin = register_and_login(client, "chefe")
    user = register_and_login(client, "comum")
    assert client.get("/auth/me", headers=admin).json()["role"] == "admin"
    assert client.get("/admin/users", headers=user).status_code == 403
    assert client.get("/admin/users", headers=admin).status_code == 200


def test_admin_lists_users_with_transaction_counts(client):
    admin = register_and_login(client, "chefe")
    user = register_and_login(client, "comum")
    add(client, user, "a")
    add(client, user, "b")
    rows = {u["username"]: u for u in client.get("/admin/users", headers=admin).json()}
    assert rows["comum"]["transaction_count"] == 2 and rows["chefe"]["transaction_count"] == 0
    assert "hashed_password" not in rows["comum"]


def test_viewer_is_read_only(client):
    admin = register_and_login(client, "chefe")
    viewer = register_and_login(client, "leitor")
    uid = client.get("/auth/me", headers=viewer).json()["id"]
    assert client.put(f"/admin/users/{uid}/role", headers=admin, json={"role": "viewer"}).status_code == 200

    body = {"description": "x", "amount": 1}
    assert client.post("/transactions/", headers=viewer, json=body).status_code == 403
    assert client.put("/budgets/lazer", headers=viewer, json={"monthly_limit": 10}).status_code == 403
    assert client.get("/transactions/", headers=viewer).status_code == 200  # leitura continua liberada
    assert client.get("/analytics/overview", headers=viewer).status_code == 200


def test_deactivated_user_cannot_login_or_use_token(client):
    admin = register_and_login(client, "chefe")
    user = register_and_login(client, "saindo")
    uid = client.get("/auth/me", headers=user).json()["id"]
    assert client.put(f"/admin/users/{uid}/active", headers=admin, json={"is_active": False}).status_code == 200
    assert client.get("/auth/me", headers=user).status_code == 401  # token antigo deixa de valer
    r = client.post("/auth/login", data={"username": "saindo", "password": "secret123"})
    assert r.status_code == 403


def test_admin_cannot_change_own_account_and_unknown_user_404(client):
    admin = register_and_login(client, "chefe")
    me = client.get("/auth/me", headers=admin).json()["id"]
    assert client.put(f"/admin/users/{me}/role", headers=admin, json={"role": "user"}).status_code == 400
    assert client.put("/admin/users/9999/active", headers=admin, json={"is_active": False}).status_code == 404


def test_overview_cache_is_invalidated_on_writes(client, auth):
    first = client.get("/analytics/overview?days=30", headers=auth).json()
    assert first["summary"]["transaction_count"] == 0
    add(client, auth, "Cafe", 10)
    second = client.get("/analytics/overview?days=30", headers=auth).json()
    assert second["summary"]["transaction_count"] == 1  # nao serviu o valor antigo do cache

    client.put("/budgets/alimentacao", headers=auth, json={"monthly_limit": 100})
    assert len(client.get("/analytics/overview?days=30", headers=auth).json()["budgets"]) == 1


def test_xlsx_export_endpoint(client, auth):
    add(client, auth, "=1+1", 10, "lazer")
    add(client, auth, "Mercado", 30, "alimentacao")
    r = client.get("/transactions/export.xlsx", headers=auth, params={"q": "mercado"})
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content))["Transações"]
    assert ws.max_row == 2 and ws["C2"].value == "Mercado"


def test_pdf_report_endpoint(client, auth):
    add(client, auth, "Mercado", 30, "alimentacao")
    r = client.get("/analytics/report.pdf?days=30", headers=auth)
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert client.get("/analytics/report.pdf").status_code == 401


def test_money_roundtrip_has_no_float_drift(client, auth):
    for _ in range(10):
        add(client, auth, "centavos", 0.1)
    s = client.get("/analytics/summary?days=30", headers=auth).json()
    assert s["total_expense"] == 1.0
