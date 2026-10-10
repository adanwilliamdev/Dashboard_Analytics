import io
from datetime import datetime, timedelta

import pandas as pd
from openpyxl import load_workbook

from app import analytics_core as core
from app import reports

NOW = datetime(2026, 10, 15, 12)


def rows():
    return [
        {"id": 1, "date": NOW - timedelta(days=3), "description": "Salario", "category": "salario", "type": "income", "amount": 5000.0},
        {"id": 2, "date": NOW - timedelta(days=2), "description": "=1+1", "category": "lazer", "type": "expense", "amount": 40.5},
        {"id": 3, "date": NOW - timedelta(days=1), "description": "Mercado", "category": "alimentacao", "type": "expense", "amount": 200.0},
    ]


def test_xlsx_structure_values_and_no_formula_injection():
    wb = load_workbook(io.BytesIO(reports.build_transactions_xlsx(rows())))
    assert wb.sheetnames == ["Resumo", "Transações", "Categorias", "Mensal"]
    ws = wb["Transações"]
    assert ws.max_row == 4 and ws["F2"].value == 5000.0
    assert ws["C3"].value == "=1+1" and ws["C3"].data_type == "s"  # texto, nao formula
    assert wb["Resumo"]["B4"].value == 4759.5  # saldo


def test_xlsx_empty():
    wb = load_workbook(io.BytesIO(reports.build_transactions_xlsx([])))
    assert wb["Transações"].max_row == 1


def make_overview():
    df = pd.DataFrame(rows())
    prev = df.iloc[0:0]
    summary = core.summary_with_comparison(df, prev, 30)
    cats = core.by_category(df)
    wd = core.weekday_profile(df)
    return {
        "days": 30,
        "summary": summary,
        "by_category": cats,
        "top_expenses": core.top_expenses(df, 5),
        "budgets": core.budget_status(df, {"alimentacao": 150.0}, today=NOW),
        "forecast": {"available": False, "reason": "x"},
        "anomalies": [],
        "insights": core.build_insights(summary, cats, wd, [], [], 30)
        + [{"level": "info", "message": "🎯 Emoji não pode quebrar o PDF"}],
    }


def test_pdf_is_valid_and_survives_emoji():
    pdf = reports.build_report_pdf(make_overview(), "ana")
    assert pdf.startswith(b"%PDF") and len(pdf) > 1500


def test_pdf_with_forecast_and_anomalies():
    ov = make_overview()
    ov["forecast"] = {"available": True, "projected_expense": 3000.0, "projected_expense_lower": 2500.0,
                      "projected_expense_upper": 3500.0, "projected_income": 5000.0, "projected_balance": 2000.0}
    ov["anomalies"] = [{"date": "2026-10-01", "description": "Compra grande", "amount": 2000.0, "expected": 100.0}]
    assert reports.build_report_pdf(ov).startswith(b"%PDF")


def test_money_format():
    assert reports.money(1234.5) == "R$ 1.234,50" and reports.money(None) == "-"
